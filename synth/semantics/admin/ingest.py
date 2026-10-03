"""
synth/semantics/admin/ingest.py

Spreadsheet and CSV ingestion engine for domain manifests (ADM-01, ADM-04).
Supports:
- .xlsx via openpyxl (read_only=True, data_only=True)
- .csv (BOM-tolerant, sniffer for ',', ';', '\\t', UTF-8 -> cp1252 fallback)
- Size caps: <= 10 MB, <= 20,000 rows
- Case/space-insensitive sheet matching and header synonym mapping
- Enum validation: ScaleType (default continuous, 'ordinal' -> categorical + warning), AxisRole (default dependent)
- Detailed RowRecord with sheet name and 1-based row number
"""
import csv
from dataclasses import dataclass, field
import io
import os
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Set, Tuple

from synth.semantics.admin.normalize import normalize_label, make_stem
from synth.semantics.enums import AxisRole, ScaleType

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB
MAX_ROW_COUNT = 20_000

# Canonical header definitions and synonyms (case/space/underscore normalized)
METRIC_HEADER_MAP: Dict[str, str] = {
    "label": "label",
    "metric": "label",
    "metriclabel": "label",
    "name": "label",
    "axislabel": "label",
    "scaletype": "scale_type",
    "scale": "scale_type",
    "type": "scale_type",
    "datatype": "scale_type",
    "axisrole": "role",
    "role": "role",
    "axis": "role",
    "unit": "unit",
    "units": "unit",
    "measurementunit": "unit",
    "conceptstem": "concept_stem",
    "stem": "concept_stem",
    "conceptgroup": "concept_group",
    "group": "concept_group",
    "domaintags": "domain_tags",
    "tags": "domain_tags",
    "domains": "domain_tags",
    "pools": "pools",
    "pool": "pools",
}

TITLE_HEADER_MAP: Dict[str, str] = {
    "title": "title",
    "charttitle": "title",
    "name": "title",
    "allowedcharttypes": "allowed_chart_types",
    "charttypes": "allowed_chart_types",
    "types": "allowed_chart_types",
    "charts": "allowed_chart_types",
    "domaintags": "domain_tags",
    "tags": "domain_tags",
    "domains": "domain_tags",
}

PAIR_HEADER_MAP: Dict[str, str] = {
    "control": "control",
    "controlgroup": "control",
    "base": "control",
    "baseline": "control",
    "treatment": "treatment",
    "treatmentgroup": "treatment",
    "intervention": "treatment",
    "test": "treatment",
    "domaincategory": "domain_category",
    "domain": "domain_category",
    "category": "domain_category",
    "contexttags": "context_tags",
    "tags": "context_tags",
}

SHEET_METRICS_NAMES = {"metrics", "metric", "axismetrics"}
SHEET_TITLES_NAMES = {"titles", "title", "charttitles"}
SHEET_PAIRS_NAMES = {"pairs", "pair", "comparativepairs"}


@dataclass(frozen=True)
class IngestIssue:
    level: str  # "error" | "warning"
    sheet: str
    row: int
    column: Optional[str]
    message: str


@dataclass(frozen=True)
class IngestedMetric:
    sheet: str
    row: int
    label: str
    scale_type: ScaleType
    role: AxisRole
    unit: Optional[str]
    concept_stem: str
    concept_group: Optional[str]
    domain_tags: Tuple[str, ...]
    pools: Tuple[str, ...]


@dataclass(frozen=True)
class IngestedTitle:
    sheet: str
    row: int
    title: str
    allowed_chart_types: Tuple[str, ...]
    domain_tags: Tuple[str, ...]


@dataclass(frozen=True)
class IngestedPair:
    sheet: str
    row: int
    control: str
    treatment: str
    domain_category: str
    context_tags: Tuple[str, ...]


@dataclass
class IngestResult:
    source_file: str
    metrics: List[IngestedMetric] = field(default_factory=list)
    titles: List[IngestedTitle] = field(default_factory=list)
    pairs: List[IngestedPair] = field(default_factory=list)
    issues: List[IngestIssue] = field(default_factory=list)

    @property
    def has_errors(self) -> bool:
        return any(i.level == "error" for i in self.issues)

    @property
    def errors(self) -> List[IngestIssue]:
        return [i for i in self.issues if i.level == "error"]

    @property
    def warnings(self) -> List[IngestIssue]:
        return [i for i in self.issues if i.level == "warning"]


def _clean_header_key(col: Any) -> str:
    """Normalize column header for synonym matching: lowercased, stripped of spaces and underscores."""
    if col is None:
        return ""
    s = str(col).lower().replace(" ", "").replace("_", "").replace("-", "")
    return s


def _parse_comma_list(val: Any) -> List[str]:
    """Parse comma, semicolon, or newline separated list of values."""
    if val is None:
        return []
    s = str(val).strip()
    if not s:
        return []
    # Split by comma or semicolon
    raw_items = [part.strip() for part in s.replace(";", ",").replace("\n", ",").split(",")]
    return [item for item in raw_items if item]


def parse_scale_type(val: Any, sheet: str, row: int, issues: List[IngestIssue]) -> ScaleType:
    """Parse ScaleType with default continuous and ordinal -> categorical mapping (ADM-04)."""
    if val is None or not str(val).strip():
        return ScaleType.CONTINUOUS

    norm = str(val).strip().lower().replace(" ", "_")
    if norm == "ordinal":
        issues.append(
            IngestIssue(
                level="warning",
                sheet=sheet,
                row=row,
                column="scale_type",
                message="'ordinal' is mapped to 'categorical'",
            )
        )
        return ScaleType.CATEGORICAL

    for member in ScaleType:
        if norm == member.value:
            return member

    allowed = [m.value for m in ScaleType] + ["ordinal"]
    issues.append(
        IngestIssue(
            level="error",
            sheet=sheet,
            row=row,
            column="scale_type",
            message=f"Invalid scale_type '{val}'. Allowed values: {allowed}",
        )
    )
    return ScaleType.CONTINUOUS


def parse_axis_role(val: Any, sheet: str, row: int, issues: List[IngestIssue]) -> AxisRole:
    """Parse AxisRole with default dependent (ADM-04)."""
    if val is None or not str(val).strip():
        return AxisRole.DEPENDENT

    norm = str(val).strip().lower().replace(" ", "_")
    for member in AxisRole:
        if norm == member.value:
            return member

    allowed = [m.value for m in AxisRole]
    issues.append(
        IngestIssue(
            level="error",
            sheet=sheet,
            row=row,
            column="role",
            message=f"Invalid axis role '{val}'. Allowed values: {allowed}",
        )
    )
    return AxisRole.DEPENDENT


def _is_excel_error_cell(val: Any) -> bool:
    if val is None:
        return False
    s = str(val).strip()
    return s in {"#N/A", "#VALUE!", "#REF!", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!"}


def ingest_file(file_path: Path | str, target_domain: str) -> IngestResult:
    """
    Ingest a domain data file (.xlsx or .csv) and parse into metrics, titles, and pairs.
    Enforces file size cap, row caps, and generates line-attributed issues.
    """
    path = Path(file_path).resolve()
    if not path.exists():
        res = IngestResult(source_file=str(file_path))
        res.issues.append(IngestIssue(level="error", sheet="", row=0, column=None, message=f"File not found: {path}"))
        return res

    file_size = path.stat().st_size
    if file_size > MAX_FILE_SIZE_BYTES:
        res = IngestResult(source_file=str(file_path))
        res.issues.append(
            IngestIssue(
                level="error",
                sheet="",
                row=0,
                column=None,
                message=f"File size {file_size / 1024 / 1024:.2f} MB exceeds maximum allowed 10 MB",
            )
        )
        return res

    ext = path.suffix.lower()
    if ext == ".xlsx":
        return _ingest_xlsx(path, target_domain)
    elif ext == ".csv":
        return _ingest_csv(path, target_domain)
    else:
        res = IngestResult(source_file=str(file_path))
        res.issues.append(
            IngestIssue(level="error", sheet="", row=0, column=None, message=f"Unsupported file format '{ext}'. Must be .xlsx or .csv")
        )
        return res


def _ingest_xlsx(path: Path, target_domain: str) -> IngestResult:
    result = IngestResult(source_file=str(path))
    import openpyxl

    try:
        wb = openpyxl.load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception as exc:
        result.issues.append(IngestIssue(level="error", sheet="", row=0, column=None, message=f"Failed to read Excel workbook: {exc}"))
        return result

    total_rows_processed = 0

    try:
        for sheet_name in wb.sheetnames:
            norm_sheet = _clean_header_key(sheet_name)
            ws = wb[sheet_name]

            # Determine sheet type
            is_metrics = norm_sheet in SHEET_METRICS_NAMES
            is_titles = norm_sheet in SHEET_TITLES_NAMES
            is_pairs = norm_sheet in SHEET_PAIRS_NAMES

            if not (is_metrics or is_titles or is_pairs):
                # If single-sheet workbook, sniff type from first row headers
                if len(wb.sheetnames) == 1:
                    is_metrics = True
                else:
                    continue

            # Read rows
            header_map: Dict[str, int] = {}
            for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                total_rows_processed += 1
                if total_rows_processed > MAX_ROW_COUNT:
                    result.issues.append(
                        IngestIssue(
                            level="error",
                            sheet=sheet_name,
                            row=row_idx,
                            column=None,
                            message=f"Workbook exceeds maximum allowed {MAX_ROW_COUNT} rows",
                        )
                    )
                    return result

                # Check for completely blank row
                if not row or all(c is None or str(c).strip() == "" for c in row):
                    continue

                # Header row
                if not header_map:
                    for col_idx, cell in enumerate(row):
                        key = _clean_header_key(cell)
                        if is_metrics and key in METRIC_HEADER_MAP:
                            header_map[METRIC_HEADER_MAP[key]] = col_idx
                        elif is_titles and key in TITLE_HEADER_MAP:
                            header_map[TITLE_HEADER_MAP[key]] = col_idx
                        elif is_pairs and key in PAIR_HEADER_MAP:
                            header_map[PAIR_HEADER_MAP[key]] = col_idx

                    # If first row didn't match headers on single-sheet, try title/pair mappings
                    if len(header_map) == 0 and len(wb.sheetnames) == 1:
                        for col_idx, cell in enumerate(row):
                            key = _clean_header_key(cell)
                            if key in TITLE_HEADER_MAP:
                                header_map[TITLE_HEADER_MAP[key]] = col_idx
                                is_titles = True
                                is_metrics = False
                        if not header_map:
                            for col_idx, cell in enumerate(row):
                                key = _clean_header_key(cell)
                                if key in PAIR_HEADER_MAP:
                                    header_map[PAIR_HEADER_MAP[key]] = col_idx
                                    is_pairs = True
                                    is_metrics = False
                    continue

                # Check for error cells in row
                error_found = False
                for cell in row:
                    if _is_excel_error_cell(cell):
                        result.issues.append(
                            IngestIssue(
                                level="error",
                                sheet=sheet_name,
                                row=row_idx,
                                column=None,
                                message=f"Excel error cell detected: '{cell}'",
                            )
                        )
                        error_found = True
                        break
                if error_found:
                    continue

                # Process data row
                if is_metrics:
                    _process_metric_row(row, header_map, sheet_name, row_idx, target_domain, result)
                elif is_titles:
                    _process_title_row(row, header_map, sheet_name, row_idx, target_domain, result)
                elif is_pairs:
                    _process_pair_row(row, header_map, sheet_name, row_idx, target_domain, result)

    finally:
        wb.close()

    return result


def _ingest_csv(path: Path, target_domain: str) -> IngestResult:
    result = IngestResult(source_file=str(path))

    # Read raw bytes with UTF-8 / cp1252 fallback
    raw_bytes = path.read_bytes()
    try:
        content = raw_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            content = raw_bytes.decode("cp1252")
        except UnicodeDecodeError as exc:
            result.issues.append(IngestIssue(level="error", sheet="csv", row=0, column=None, message=f"Encoding error: {exc}"))
            return result

    # Sniff delimiter
    first_lines = "\n".join(content.splitlines()[:10])
    try:
        dialect = csv.Sniffer().sniff(first_lines, delimiters=",\t;")
        delimiter = dialect.delimiter
    except Exception:
        delimiter = ","

    reader = csv.reader(io.StringIO(content), delimiter=delimiter)
    header_map: Dict[str, int] = {}
    sheet_type = "metrics"  # default

    for row_idx, row in enumerate(reader, start=1):
        if row_idx > MAX_ROW_COUNT:
            result.issues.append(
                IngestIssue(level="error", sheet="csv", row=row_idx, column=None, message=f"CSV exceeds maximum allowed {MAX_ROW_COUNT} rows")
            )
            return result

        if not row or all(c.strip() == "" for c in row):
            continue

        if not header_map:
            # Match headers
            m_matches = sum(1 for c in row if _clean_header_key(c) in METRIC_HEADER_MAP)
            t_matches = sum(1 for c in row if _clean_header_key(c) in TITLE_HEADER_MAP)
            p_matches = sum(1 for c in row if _clean_header_key(c) in PAIR_HEADER_MAP)

            if t_matches > m_matches and t_matches > p_matches:
                sheet_type = "titles"
                h_map = TITLE_HEADER_MAP
            elif p_matches > m_matches and p_matches > t_matches:
                sheet_type = "pairs"
                h_map = PAIR_HEADER_MAP
            else:
                sheet_type = "metrics"
                h_map = METRIC_HEADER_MAP

            for col_idx, cell in enumerate(row):
                key = _clean_header_key(cell)
                if key in h_map:
                    header_map[h_map[key]] = col_idx
            continue

        if sheet_type == "metrics":
            _process_metric_row(row, header_map, "csv", row_idx, target_domain, result)
        elif sheet_type == "titles":
            _process_title_row(row, header_map, "csv", row_idx, target_domain, result)
        elif sheet_type == "pairs":
            _process_pair_row(row, header_map, "csv", row_idx, target_domain, result)

    return result


def _get_val(row: Sequence[Any], header_map: Dict[str, int], key: str) -> Optional[Any]:
    idx = header_map.get(key)
    if idx is not None and idx < len(row):
        val = row[idx]
        if val is not None and str(val).strip() != "":
            return val
    return None


def _process_metric_row(
    row: Sequence[Any],
    header_map: Dict[str, int],
    sheet: str,
    row_idx: int,
    target_domain: str,
    result: IngestResult,
):
    raw_label = _get_val(row, header_map, "label")
    if not raw_label:
        result.issues.append(IngestIssue(level="error", sheet=sheet, row=row_idx, column="label", message="Metric label cannot be empty"))
        return

    label = normalize_label(str(raw_label))
    scale_type = parse_scale_type(_get_val(row, header_map, "scale_type"), sheet, row_idx, result.issues)
    role = parse_axis_role(_get_val(row, header_map, "role"), sheet, row_idx, result.issues)

    unit_val = _get_val(row, header_map, "unit")
    unit = normalize_label(str(unit_val)) if unit_val else None

    stem_val = _get_val(row, header_map, "concept_stem")
    if stem_val:
        stem = make_stem(str(stem_val))
    else:
        stem = make_stem(label)

    if not stem:
        result.issues.append(
            IngestIssue(level="error", sheet=sheet, row=row_idx, column="concept_stem", message=f"Failed to generate valid concept_stem for label '{label}'")
        )
        return

    grp_val = _get_val(row, header_map, "concept_group")
    concept_group = make_stem(str(grp_val)) if grp_val else None

    # Domain tags: always ensure target_domain is included (H1)
    tags_val = _get_val(row, header_map, "domain_tags")
    parsed_tags = _parse_comma_list(tags_val)
    tags = [target_domain] + [t for t in parsed_tags if t != target_domain]

    # Pools
    pools_val = _get_val(row, header_map, "pools")
    pools = _parse_comma_list(pools_val)

    metric = IngestedMetric(
        sheet=sheet,
        row=row_idx,
        label=label,
        scale_type=scale_type,
        role=role,
        unit=unit,
        concept_stem=stem,
        concept_group=concept_group,
        domain_tags=tuple(tags),
        pools=tuple(pools),
    )
    result.metrics.append(metric)


def _process_title_row(
    row: Sequence[Any],
    header_map: Dict[str, int],
    sheet: str,
    row_idx: int,
    target_domain: str,
    result: IngestResult,
):
    raw_title = _get_val(row, header_map, "title")
    if not raw_title:
        result.issues.append(IngestIssue(level="error", sheet=sheet, row=row_idx, column="title", message="Title cannot be empty"))
        return

    title = normalize_label(str(raw_title))

    # Allowed chart types: default to all cartesian types if blank (ADM-05)
    types_val = _get_val(row, header_map, "allowed_chart_types")
    if types_val:
        chart_types = [t.lower().strip() for t in _parse_comma_list(types_val)]
    else:
        from synth.semantics.loader import CARTESIAN_CHART_TYPES
        chart_types = list(CARTESIAN_CHART_TYPES)

    # Validate chart types
    from synth.semantics.loader import ALL_CHART_TYPES
    for ct in chart_types:
        if ct not in ALL_CHART_TYPES:
            result.issues.append(
                IngestIssue(
                    level="error",
                    sheet=sheet,
                    row=row_idx,
                    column="allowed_chart_types",
                    message=f"Invalid chart type '{ct}'. Allowed: {sorted(ALL_CHART_TYPES)}",
                )
            )

    tags_val = _get_val(row, header_map, "domain_tags")
    parsed_tags = _parse_comma_list(tags_val)
    tags = [target_domain] + [t for t in parsed_tags if t != target_domain]

    title_obj = IngestedTitle(
        sheet=sheet,
        row=row_idx,
        title=title,
        allowed_chart_types=tuple(chart_types),
        domain_tags=tuple(tags),
    )
    result.titles.append(title_obj)


def _process_pair_row(
    row: Sequence[Any],
    header_map: Dict[str, int],
    sheet: str,
    row_idx: int,
    target_domain: str,
    result: IngestResult,
):
    raw_ctrl = _get_val(row, header_map, "control")
    raw_treat = _get_val(row, header_map, "treatment")

    if not raw_ctrl or not raw_treat:
        result.issues.append(
            IngestIssue(level="error", sheet=sheet, row=row_idx, column=None, message="Pair control and treatment cannot be empty")
        )
        return

    control = normalize_label(str(raw_ctrl))
    treatment = normalize_label(str(raw_treat))

    if control == treatment:
        result.issues.append(
            IngestIssue(
                level="error",
                sheet=sheet,
                row=row_idx,
                column=None,
                message=f"Pair control and treatment must be distinct (got '{control}')",
            )
        )
        return

    cat_val = _get_val(row, header_map, "domain_category")
    category = normalize_label(str(cat_val)) if cat_val else target_domain

    tags_val = _get_val(row, header_map, "context_tags")
    parsed_tags = _parse_comma_list(tags_val)
    tags = [category] + [t for t in parsed_tags if t != category]

    pair_obj = IngestedPair(
        sheet=sheet,
        row=row_idx,
        control=control,
        treatment=treatment,
        domain_category=category,
        context_tags=tuple(tags),
    )
    result.pairs.append(pair_obj)
