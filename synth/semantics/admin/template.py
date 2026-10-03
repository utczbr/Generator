"""
synth/semantics/admin/template.py

Excel domain spreadsheet template generator (ADM-14).
Generates domain_template.xlsx containing:
1. 'Metrics' sheet with exact headers and Scale Type / Axis Role DataValidation dropdowns
2. 'Titles' sheet with exact headers and Cartesian / special chart types documentation
3. 'Comparative Pairs' sheet with exact headers and example row
"""
from pathlib import Path
from typing import Optional, Union
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

from synth.semantics.enums import AxisRole, ScaleType
from synth.semantics.loader import ALL_CHART_TYPES


def generate_template(output_path: Union[str, Path] = "domain_template.xlsx") -> Path:
    """
    Generate an annotated Excel template workbook for authoring domain manifests (ADM-14).
    """
    out_file = Path(output_path).resolve()
    wb = openpyxl.Workbook()

    # Define styles
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # --------------------------------------------------------------------------
    # 1. Sheet: Metrics
    # --------------------------------------------------------------------------
    ws_metrics = wb.active
    ws_metrics.title = "Metrics"

    metrics_headers = [
        "Label",
        "Scale Type",
        "Axis Role",
        "Unit",
        "Concept Stem",
        "Concept Group",
        "Domain Tags",
        "Pools",
    ]
    ws_metrics.append(metrics_headers)

    # Example row
    ws_metrics.append([
        "Core Rotor Temp (°C)",
        "continuous",
        "dependent",
        "°C",
        "core_rotor_temp",
        "temperature_metric",
        "aerospace",
        "legacy",
    ])

    # Header notes
    ws_metrics["A1"].comment = Comment("Human-facing axis label. Superscripts (m²) and symbols (μM) preserved.", "System")
    ws_metrics["B1"].comment = Comment("Must be one of: continuous, categorical, discrete_count, percentage, temporal.", "System")
    ws_metrics["C1"].comment = Comment("Must be one of: independent, dependent, bidirectional.", "System")
    ws_metrics["E1"].comment = Comment("Identifier format ^[a-z0-9_]+$. If omitted, auto-generated from Label.", "System")

    # Dropdowns for Scale Type and Axis Role
    scale_types_str = f'"{",".join(s.value for s in ScaleType)}"'
    dv_scale = DataValidation(type="list", formula1=scale_types_str, allow_blank=True)
    ws_metrics.add_data_validation(dv_scale)
    dv_scale.add("B2:B2000")

    roles_str = f'"{",".join(r.value for r in AxisRole)}"'
    dv_role = DataValidation(type="list", formula1=roles_str, allow_blank=True)
    ws_metrics.add_data_validation(dv_role)
    dv_role.add("C2:C2000")

    # Column widths
    metric_widths = [26, 16, 16, 10, 22, 22, 16, 14]
    for idx, width in enumerate(metric_widths, start=1):
        col_letter = openpyxl.utils.get_column_letter(idx)
        ws_metrics.column_dimensions[col_letter].width = width

    # --------------------------------------------------------------------------
    # 2. Sheet: Titles
    # --------------------------------------------------------------------------
    ws_titles = wb.create_sheet(title="Titles")
    titles_headers = ["Title", "Allowed Chart Types", "Domain Tags"]
    ws_titles.append(titles_headers)

    # Example row
    ws_titles.append([
        "Turbine Blade Thermal Distribution",
        "bar, line, scatter, box, area, histogram",
        "aerospace",
    ])

    ws_titles["B1"].comment = Comment(
        f"Comma-separated list of chart types.\nAllowed values: {', '.join(sorted(ALL_CHART_TYPES))}.\nLeave blank for all Cartesian types.",
        "System",
    )

    title_widths = [36, 42, 18]
    for idx, width in enumerate(title_widths, start=1):
        col_letter = openpyxl.utils.get_column_letter(idx)
        ws_titles.column_dimensions[col_letter].width = width

    # --------------------------------------------------------------------------
    # 3. Sheet: Comparative Pairs
    # --------------------------------------------------------------------------
    ws_pairs = wb.create_sheet(title="Comparative Pairs")
    pairs_headers = ["Control", "Treatment", "Domain Category", "Context Tags"]
    ws_pairs.append(pairs_headers)

    # Example row
    ws_pairs.append([
        "Baseline Alloy",
        "Titanium Composite",
        "aerospace",
        "aerospace, materials",
    ])

    ws_pairs["A1"].comment = Comment("Baseline / Control group label.", "System")
    ws_pairs["B1"].comment = Comment("Treatment / Intervention group label (must be distinct from Control).", "System")

    pair_widths = [24, 24, 20, 24]
    for idx, width in enumerate(pair_widths, start=1):
        col_letter = openpyxl.utils.get_column_letter(idx)
        ws_pairs.column_dimensions[col_letter].width = width

    # Format headers across all sheets
    for ws in (ws_metrics, ws_titles, ws_pairs):
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.font = data_font
                cell.border = thin_border

    out_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_file))
    return out_file
